#!/usr/bin/env python3
"""Rebuilds the state data of the site from the official files kept in research/ (rejouable).

  python3 scripts/build-data.py

Inputs (research/):
  usda-bbce-states-june2026.txt        USDA FNA BBCE chart (page updated June 29, 2026)
  usda-sua-table-fy2026.xlsx           USDA FNA SUA table FY 2026 (May 21, 2026)
  cms-medicaid-chip-eligibility-levels.txt  Medicaid.gov table (as of December 1, 2023)
  hud-section8-income-limits-fy2026.xlsx    HUD FY 2026 Section 8 income limits (May 1, 2026)
  state-facts.json                     state agencies, program names, marketplaces (official URLs checked)
Outputs:
  src/data/states.json                 one object per state (50 + DC)
  src/data/hud/<ST>.json               every HUD area/county row of the state
"""
import json, os, re, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R = lambda f: os.path.join(ROOT, 'research', f)
OUT = os.path.join(ROOT, 'src', 'data')

STATES = [
    ('alabama', 'Alabama', 'AL'), ('alaska', 'Alaska', 'AK'), ('arizona', 'Arizona', 'AZ'), ('arkansas', 'Arkansas', 'AR'),
    ('california', 'California', 'CA'), ('colorado', 'Colorado', 'CO'), ('connecticut', 'Connecticut', 'CT'), ('delaware', 'Delaware', 'DE'),
    ('district-of-columbia', 'District of Columbia', 'DC'), ('florida', 'Florida', 'FL'), ('georgia', 'Georgia', 'GA'), ('hawaii', 'Hawaii', 'HI'),
    ('idaho', 'Idaho', 'ID'), ('illinois', 'Illinois', 'IL'), ('indiana', 'Indiana', 'IN'), ('iowa', 'Iowa', 'IA'),
    ('kansas', 'Kansas', 'KS'), ('kentucky', 'Kentucky', 'KY'), ('louisiana', 'Louisiana', 'LA'), ('maine', 'Maine', 'ME'),
    ('maryland', 'Maryland', 'MD'), ('massachusetts', 'Massachusetts', 'MA'), ('michigan', 'Michigan', 'MI'), ('minnesota', 'Minnesota', 'MN'),
    ('mississippi', 'Mississippi', 'MS'), ('missouri', 'Missouri', 'MO'), ('montana', 'Montana', 'MT'), ('nebraska', 'Nebraska', 'NE'),
    ('nevada', 'Nevada', 'NV'), ('new-hampshire', 'New Hampshire', 'NH'), ('new-jersey', 'New Jersey', 'NJ'), ('new-mexico', 'New Mexico', 'NM'),
    ('new-york', 'New York', 'NY'), ('north-carolina', 'North Carolina', 'NC'), ('north-dakota', 'North Dakota', 'ND'), ('ohio', 'Ohio', 'OH'),
    ('oklahoma', 'Oklahoma', 'OK'), ('oregon', 'Oregon', 'OR'), ('pennsylvania', 'Pennsylvania', 'PA'), ('rhode-island', 'Rhode Island', 'RI'),
    ('south-carolina', 'South Carolina', 'SC'), ('south-dakota', 'South Dakota', 'SD'), ('tennessee', 'Tennessee', 'TN'), ('texas', 'Texas', 'TX'),
    ('utah', 'Utah', 'UT'), ('vermont', 'Vermont', 'VT'), ('virginia', 'Virginia', 'VA'), ('washington', 'Washington', 'WA'),
    ('west-virginia', 'West Virginia', 'WV'), ('wisconsin', 'Wisconsin', 'WI'), ('wyoming', 'Wyoming', 'WY'),
]
NAMES = {n: (s, a) for s, n, a in STATES}

# ---------------------------------------------------------------- BBCE
def bbce():
    out = {}
    for line in open(R('usda-bbce-states-june2026.txt')):
        cells = [c.strip() for c in line.split('|')]
        if len(cells) < 4 or cells[0] not in NAMES:
            continue
        name, desc, asset, gross = cells[0], cells[1], cells[2], cells[3]
        m = re.match(r'(\d+)%', gross)
        if not m:
            continue
        pct = int(m.group(1))
        lim = None if asset.startswith('No limit') else int(re.search(r'\$([\d,]+)', asset).group(1).replace(',', ''))
        who = 'all'
        if 'non-elderly' in desc: who = 'non-elderly-disabled'
        elif 'elderly or disabled' in desc: who = 'elderly-disabled'
        elif 'dependent care' in desc: who = 'dependent-care'
        elif 'earned income' in desc: who = 'earned-income'
        out.setdefault(name, []).append({'who': who, 'grossPct': pct, 'assetLimit': lim, 'assetText': asset, 'program': desc})
    return out

# ---------------------------------------------------------------- SUA
def sua():
    wb = openpyxl.load_workbook(R('usda-sua-table-fy2026.xlsx'), data_only=True)
    ws = wb.active
    out = {}
    for row in list(ws.iter_rows(values_only=True))[3:]:
        label = row[0]
        if not label or not isinstance(label, str):
            continue
        base = re.split(r'\s+-\s+', label.replace('*', '').strip(), maxsplit=1)
        name = base[0].strip()
        if name not in NAMES:
            continue
        variant = base[1].strip() if len(base) > 1 else None
        vals = row[1:10]
        if vals[0] == '-':
            out.setdefault(name, []).append({'variant': variant, 'pending': True})
            continue
        num = lambda x: (round(float(x), 2) if x not in (None, '-') else 0)
        out.setdefault(name, []).append({
            'variant': variant, 'hcsua': num(vals[0]), 'lua': num(vals[1]), 'electricity': num(vals[2]), 'gas': num(vals[3]),
            'water': num(vals[4]), 'sewage': num(vals[5]), 'trash': num(vals[6]), 'phone': num(vals[7]), 'smdOffset': num(vals[8]),
            'mandatory': '*' not in label,
        })
    return out

# ---------------------------------------------------------------- Medicaid (CMS)
def pctval(s):
    m = re.match(r'(\d+)%', s.strip())
    return int(m.group(1)) if m else None

def medicaid():
    out = {}
    for line in open(R('cms-medicaid-chip-eligibility-levels.txt')):
        cells = [c.strip() for c in line.split('|')]
        if len(cells) < 9:
            continue
        name = re.sub(r'\s+\d+$', '', cells[0]).strip()
        if name not in NAMES:
            continue
        c01, c15, c618, chip, pw, pwchip, parent, adults = cells[1:9]
        chip_pct = pctval(chip)
        chip_ages = None
        m = re.search(r'\(([^)]*)\)', chip)
        if m: chip_ages = m.group(1).replace('up to19', 'up to 19')
        note_refs = re.findall(r'(?<=\s)(\d{1,2})(?=\s|$)', ' '.join([c01, chip, adults, parent]))
        expansion = None; adult_pct = None; bhp = None
        a = adults.strip()
        if a.startswith('No/'):
            expansion = False; adult_pct = pctval(a[3:])
        elif a.startswith('No'):
            expansion = False
        else:
            expansion = True; adult_pct = pctval(a)
            if re.search(r'\b11$', a):  # adults above 133% to 200% through the Basic Health Program
                bhp = adult_pct; adult_pct = 133
        par = pctval(parent)
        if name == 'Wisconsin':
            par = pctval(parent)
        out[name] = {
            'child0to1': pctval(c01), 'child1to5': pctval(c15), 'child6to18': pctval(c618),
            'chip': chip_pct, 'chipAges': chip_ages, 'chipNote': None if chip_pct is not None or chip.startswith('N/A') else chip,
            'pregnant': pctval(pw), 'pregnantChip': pctval(pwchip),
            'parent': par, 'parentDollarStandard': '($)' in parent,
            'expansion': expansion, 'adultPct': adult_pct, 'bhpPct': bhp,
            'raw': {'adults': adults, 'parent': parent, 'chip': chip},
        }
    return out

# Footnotes of the CMS table that change the reading of a cell (text of research/cms-medicaid-chip-eligibility-levels.txt).
CMS_NOTES = {
    'California': 'CMS note 6: the separate CHIP in California covers certain children up to age 2 with incomes up to 317% of the FPL statewide, and children up to age 19 up to 317% in three counties only.',
    'District of Columbia': 'CMS note 7: DC also covers ages 19 and 20 up to 216% of the FPL.',
    'Massachusetts': 'CMS note 7: Massachusetts also covers ages 19 and 20 up to 150% of the FPL.',
    'Maine': 'CMS note 7: Maine also covers ages 19 and 20 up to 300% of the FPL.',
    'Florida': 'CMS note 8: Florida also covers ages 19 and 20 up to 24% of the FPL.',
    'Minnesota': 'CMS notes 10 and 11: Minnesota covers children up to age 2 to 283% of the FPL; adults above 133% and up to 200% are covered through the Basic Health Program (MinnesotaCare).',
    'New York': 'CMS note 11: adults above 133% and up to 200% of the FPL are covered through the Basic Health Program.',
    'Oklahoma': 'CMS note 9: section 1115 demonstration covering some additional low-income adults.',
    'South Carolina': 'CMS note 9: section 1115 demonstration covering some additional low-income adults, with limits.',
    'Utah': 'CMS note 9: section 1115 demonstration covering some additional low-income adults.',
    'Wisconsin': 'CMS note 9: Wisconsin did not adopt the ACA expansion; a section 1115 demonstration covers adults up to 95% of the FPL (100% with the 5% disregard).',
}

# ---------------------------------------------------------------- HUD FY 2026
def hud():
    wb = openpyxl.load_workbook(R('hud-section8-income-limits-fy2026.xlsx'), read_only=True, data_only=True)
    ws = wb['Section8-FY26']
    rows = list(ws.iter_rows(values_only=True)); hdr = rows[0]
    by = {}
    for r in rows[1:]:
        d = dict(zip(hdr, r))
        st = d['stusps']
        if st not in {a for _, _, a in STATES}:
            continue
        # New England rows are towns (county_town_name); elsewhere the county name is complete,
        # while county_town_name is cut at 26 characters ("Fairbanks North Star Borou").
        town = st in ('CT', 'ME', 'MA', 'NH', 'RI', 'VT')
        name = ((d['county_town_name'] if town else d['County_Name']) or d['County_Name']).strip()
        by.setdefault(st, []).append({
            'n': name, 'county': d['County_Name'].strip(), 'area': d['hud_area_name'].strip(), 'metro': int(d['metro'] or 0), 'median': int(d['median2026']),
            'v': [int(d[f'l50_{i}']) for i in range(1, 9)], 'e': [int(d[f'ELI_{i}']) for i in range(1, 9)], 'l': [int(d[f'l80_{i}']) for i in range(1, 9)],
        })
    return by

def main():
    B, S, M, H = bbce(), sua(), medicaid(), hud()
    # Changes since the CMS table, each proved by the state's own official page (research/medicaid-updates.json).
    upd = {}
    up = R('medicaid-updates.json')
    if os.path.exists(up):
        upd = {u['slug']: u for u in json.load(open(up)) if u.get('status') == 'changed' and u.get('changes') and u.get('sources')}
    facts = {}
    fp = R('state-facts.json')
    if os.path.exists(fp):
        facts = {f['name']: f for f in json.load(open(fp))}
    os.makedirs(os.path.join(OUT, 'hud'), exist_ok=True)
    # Two rows with the same town name in one state (Maine) are told apart by their county.
    for st, rows in H.items():
        seen = {}
        for r in rows: seen[r['n']] = seen.get(r['n'], 0) + 1
        for r in rows:
            if seen[r['n']] > 1: r['n'] = f"{r['n']} ({r['county']})"

    states = []
    for slug, name, abbr in STATES:
        rows = sorted(H[abbr], key=lambda x: x['n'])
        json.dump({'abbr': abbr, 'rows': rows}, open(os.path.join(OUT, 'hud', f'{abbr}.json'), 'w'), separators=(',', ':'))
        v4 = [r['v'][3] for r in rows]
        hi = max(rows, key=lambda r: r['v'][3]); lo = min(rows, key=lambda r: r['v'][3])
        f = facts.get(name, {})
        med = dict(M[name]); med.pop('raw', None)
        med['note'] = CMS_NOTES.get(name)
        u = upd.get(slug)
        if u:
            for k, v in u['changes'].items():
                if k in med: med[k] = v
            med['update'] = {'effective': u.get('effective'), 'note': u.get('note'), 'sources': [{'url': x['url'], 'title': x.get('title', '')} for x in u['sources']]}
            if med['note'] and slug == 'district-of-columbia': med['note'] = None
        else:
            med['update'] = None
        states.append({
            'slug': slug, 'name': name, 'abbr': abbr,
            'area': 'ak' if abbr == 'AK' else 'hi' if abbr == 'HI' else 'us',
            'snap': {'bbce': B.get(name, []), 'sua': S.get(name, []), **{k: v for k, v in (f.get('snap') or {}).items() if k != 'urlStatus'}},
            'medicaid': {**med, **{k: v for k, v in (f.get('medicaid') or {}).items() if k != 'urlStatus'}},
            'marketplace': {k: v for k, v in (f.get('marketplace') or {}).items() if k != 'urlStatus'} or None,
            'wic': {k: v for k, v in (f.get('wic') or {}).items() if k != 'urlStatus'} or None,
            'notes': f.get('notes', []),
            'hud': {'rows': len(rows), 'areas': len({r['area'] for r in rows}), 'vli4Max': hi['v'][3], 'vli4MaxName': hi['n'], 'vli4MaxArea': hi['area'],
                    'vli4Min': lo['v'][3], 'vli4MinName': lo['n'], 'vli4MinArea': lo['area']},
        })
    json.dump(states, open(os.path.join(OUT, 'states.json'), 'w'), indent=1, ensure_ascii=False)
    print(f"Medicaid updates applied: {sorted(upd)}")
    print(f"{len(states)} states · BBCE {sum(1 for s in states if s['snap']['bbce'])} · SUA {sum(1 for s in states if s['snap']['sua'])} · HUD rows {sum(s['hud']['rows'] for s in states)} · facts {len(facts)}")

if __name__ == '__main__':
    main()
